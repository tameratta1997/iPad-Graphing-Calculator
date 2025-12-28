from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.db.models import Sum, F, Count
from django.db import models
from django.contrib import messages
from django.contrib.auth.models import User
from django.contrib.auth.forms import UserCreationForm
from django.contrib.auth.decorators import login_required, user_passes_test, permission_required
import csv
import openpyxl
from .models import Product, StockLog, Category
from .forms import ProductForm, ImportFileForm, CustomUserCreationForm, CustomUserChangeForm

@login_required
def dashboard(request):
    total_products = Product.objects.count()
    total_stock_value = Product.objects.aggregate(
        total_value=Sum(F('quantity') * F('purchase_price'))
    )['total_value'] or 0
    
    low_stock_products = Product.objects.filter(quantity__lte=F('min_stock_alert'))
    recent_activity = StockLog.objects.select_related('product', 'user').order_by('-timestamp')[:5]

    context = {
        'total_products': total_products,
        'total_stock_value': total_stock_value,
        'low_stock_count': low_stock_products.count(),
        'recent_activity': recent_activity,
        'low_stock_products': low_stock_products,
    }
    return render(request, 'inventory/dashboard.html', context)

@login_required
def product_list(request):
    products = Product.objects.select_related('category').all()
    query = request.GET.get('q')
    if query:
        products = products.filter(name__icontains=query) | products.filter(code__icontains=query)
    
    context = {'products': products}
    return render(request, 'inventory/product_list.html', context)

@login_required
@permission_required('inventory.add_product', raise_exception=True)
def product_create(request):
    if request.method == 'POST':
        form = ProductForm(request.POST, request.FILES)
        if form.is_valid():
            product = form.save()
            StockLog.objects.create(
                product=product,
                user=request.user,
                action='ADD',
                quantity_change=product.quantity,
                reason='Initial Creation'
            )
            return redirect('product_list')
    else:
        form = ProductForm()
    return render(request, 'inventory/product_form.html', {'form': form, 'title': 'Add Product'})

@login_required
@permission_required('inventory.change_product', raise_exception=True)
def product_update(request, pk):
    product = get_object_or_404(Product, pk=pk)
    old_qty = product.quantity
    if request.method == 'POST':
        form = ProductForm(request.POST, request.FILES, instance=product)
        if form.is_valid():
            product = form.save()
            # Log stock change if any
            if product.quantity != old_qty:
                change = product.quantity - old_qty
                StockLog.objects.create(
                    product=product,
                    user=request.user,
                    action='ADJUST',
                    quantity_change=change,
                    reason='Update via Form'
                )
            return redirect('product_list')
    else:
        form = ProductForm(instance=product)
    return render(request, 'inventory/product_form.html', {'form': form, 'title': 'Edit Product'})

@login_required
@permission_required('inventory.delete_product', raise_exception=True)
def product_delete(request, pk):
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        product.delete()
        return redirect('product_list')
    return render(request, 'inventory/product_confirm_delete.html', {'product': product})

@login_required
def category_list(request):
    categories = Category.objects.annotate(product_count=Count('products')).all()
    return render(request, 'inventory/category_list.html', {'categories': categories})

@login_required
@permission_required('inventory.add_category', raise_exception=True)
def category_create(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        description = request.POST.get('description')
        if name:
            Category.objects.create(name=name, description=description)
            return redirect('category_list')
    return render(request, 'inventory/category_form.html')

@login_required
@permission_required('inventory.delete_category', raise_exception=True)
def category_delete(request, pk):
    category = get_object_or_404(Category, pk=pk)
    if request.method == 'POST':
        category.delete()
        return redirect('category_list')
    return render(request, 'inventory/category_confirm_delete.html', {'category': category})

@login_required
@permission_required('inventory.add_product', raise_exception=True)
def product_import(request):
    if request.method == 'POST':
        form = ImportFileForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            created_count = 0
            errors = []
            
            try:
                # Identify file type
                is_excel = file.name.endswith(('.xlsx', '.xls'))
                
                rows = []
                if is_excel:
                    wb = openpyxl.load_workbook(file)
                    ws = wb.active
                    # Skip header
                    iter_rows = ws.iter_rows(values_only=True)
                    next(iter_rows, None) 
                    rows = list(iter_rows)
                else:
                    # CSV handling
                    decoded_file = file.read().decode('utf-8').splitlines()
                    reader = csv.reader(decoded_file)
                    next(reader, None) # Skip header
                    rows = list(reader)

                for row in rows:
                    try:
                        # Expected format: Name, Code, Category, Purchase Price, Selling Price, Qty, Min Stock
                        # Adjust index checking based on format. Assuming standard 7 columns for now.
                        if not row or len(row) < 3: continue 
                        
                        name = row[0]
                        code = str(row[1])
                        category_name = row[2]
                        p_price = row[3] if len(row) > 3 else 0
                        s_price = row[4] if len(row) > 4 else 0
                        qty = row[5] if len(row) > 5 else 0
                        min_alert = row[6] if len(row) > 6 else 10

                        # Get or Create Category
                        category, _ = Category.objects.get_or_create(name=category_name)
                        
                        # Create Product (avoid duplicates by code)
                        if not Product.objects.filter(code=code).exists():
                            product = Product.objects.create(
                                name=name,
                                code=code,
                                category=category,
                                purchase_price=p_price,
                                selling_price=s_price,
                                quantity=qty,
                                min_stock_alert=min_alert
                            )
                            created_count += 1
                            
                            # Log initial stock
                            StockLog.objects.create(
                                product=product,
                                user=request.user,
                                action='ADD',
                                quantity_change=qty,
                                reason='Bulk Import'
                            )
                    except Exception as e:
                        errors.append(f"Error row {row}: {str(e)}")
                
                if created_count > 0:
                    messages.success(request, f"Successfully imported {created_count} products.")
                else:
                    messages.warning(request, "No new products imported. Check codes for duplicates.")
                    
                if errors:
                    for err in errors[:5]: # Show first 5 errors
                        messages.error(request, err)
                        
                return redirect('product_list')

            except Exception as e:
                messages.error(request, f"Critical Error: {str(e)}")

    else:
        form = ImportFileForm()
    
    return render(request, 'inventory/product_import.html', {'form': form})

@user_passes_test(lambda u: u.is_superuser)
def user_list(request):
    users = User.objects.all()
    return render(request, 'inventory/user_list.html', {'users': users})

@user_passes_test(lambda u: u.is_superuser)
def user_create(request):
    if request.method == 'POST':
        form = CustomUserCreationForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'New user created and permissions assigned successfully.')
            return redirect('user_list')
    else:
        form = CustomUserCreationForm()
    return render(request, 'inventory/user_form.html', {'form': form, 'title': 'Create New User'})

@user_passes_test(lambda u: u.is_superuser)
def user_update(request, pk):
    user_to_edit = get_object_or_404(User, pk=pk)
    if request.method == 'POST':
        form = CustomUserChangeForm(request.POST, instance=user_to_edit)
        if form.is_valid():
            form.save()
            messages.success(request, f'User {user_to_edit.username} updated successfully.')
            return redirect('user_list')
    else:
        form = CustomUserChangeForm(instance=user_to_edit)
    return render(request, 'inventory/user_form.html', {'form': form, 'title': f'Edit User: {user_to_edit.username}'})

@user_passes_test(lambda u: u.is_superuser)
def user_delete(request, pk):
    user_to_delete = get_object_or_404(User, pk=pk)
    
    # Prevent self-deletion
    if user_to_delete == request.user:
        messages.error(request, "You cannot delete your own account.")
        return redirect('user_list')
        
    if request.method == 'POST':
        user_to_delete.delete()
        messages.success(request, f"User {user_to_delete.username} deleted successfully.")
        return redirect('user_list')
    
    return render(request, 'inventory/user_confirm_delete.html', {'user_to_delete': user_to_delete})

# API Views
from rest_framework import viewsets
from .serializers import ProductSerializer, CategorySerializer
from .models import Category

class ProductViewSet(viewsets.ModelViewSet):
    queryset = Product.objects.all()
    serializer_class = ProductSerializer

class CategoryViewSet(viewsets.ModelViewSet):
    queryset = Category.objects.all()
    serializer_class = CategorySerializer

